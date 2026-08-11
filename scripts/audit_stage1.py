#!/usr/bin/env python3
"""Read-only, streaming inventory for the stage-1 source data audit.

The script never writes below ``data/``.  CSV/TSV files are streamed so the
largest tables do not need to fit in memory.  Field roles are deliberately
reported as *candidates*: label semantics still require human registration.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Sequence


KNOWN_HEADER_FIELDS = {
    "heavy", "light", "fitness", "sequence", "pred_affinity", "target",
    "ab_name", "ag_name", "antigen", "id", "chain",
}
SEQUENCE_ROLE_PATTERNS = {
    "heavy": (r"^heavy$", r"heavy_chain", r"^h_seq$", r"ab_heavy", r"ab_nano_h_chain_aa"),
    "light": (r"^light$", r"light_chain", r"^l_seq$", r"ab_light", r"ab_nano_l_chain_aa"),
    "antigen": (r"antigen_seq", r"^ag_seq$", r"target_seq"),
}
LABEL_HINTS = (
    "fitness", "affinity", "kd", "ic50", "ec50", "binding", "neutralization",
    "log_aff", "aff_op",
)
AA_RE = re.compile(r"^[ACDEFGHIKLMNPQRSTVWYBXZJUO*-]+$", re.I)


@dataclass
class InventoryRow:
    path: str
    format: str
    size_bytes: int
    sha256: str
    sheet: str
    header_row: int | None
    row_count: int | None
    column_count: int | None
    bad_width_rows: int | None
    columns: list[str]
    heavy_candidates: list[str]
    light_candidates: list[str]
    antigen_candidates: list[str]
    label_candidates: list[str]
    sampled_rows: int
    sampled_missing: dict[str, int]
    sampled_invalid_sequence: dict[str, int]
    review_status: str
    notes: list[str]


def normalize(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_")


def sha256_file(path: Path, chunk_size: int = 4 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def looks_like_header(row: Sequence[str]) -> bool:
    normalized = {normalize(value) for value in row if value.strip()}
    known_hits = len(normalized & KNOWN_HEADER_FIELDS)
    role_hits = sum(
        any(re.search(pattern, value) for pattern in sum(SEQUENCE_ROLE_PATTERNS.values(), ()))
        for value in normalized
    )
    return known_hits >= 1 and (known_hits + role_hits) >= 2


def infer_roles(columns: Sequence[str]) -> dict[str, list[str]]:
    roles = {"heavy": [], "light": [], "antigen": [], "label": []}
    for original in columns:
        value = normalize(original)
        for role, patterns in SEQUENCE_ROLE_PATTERNS.items():
            if any(re.search(pattern, value) for pattern in patterns):
                roles[role].append(original)
        if any(hint in value for hint in LABEL_HINTS):
            roles["label"].append(original)
    # A generic sequence column is a plausible VHH/protein candidate only when
    # no explicit heavy-chain field exists.  It remains a review candidate.
    if not roles["heavy"]:
        roles["heavy"] = [name for name in columns if normalize(name) == "sequence"]
    return roles


def _sample_stats(
    columns: Sequence[str], rows: Iterable[Sequence[object]], roles: dict[str, list[str]], limit: int
) -> tuple[int, dict[str, int], dict[str, int]]:
    indices = {name: idx for idx, name in enumerate(columns)}
    relevant = list(dict.fromkeys(sum(roles.values(), [])))
    missing = Counter({name: 0 for name in relevant})
    invalid = Counter({name: 0 for name in sum((roles[r] for r in ("heavy", "light", "antigen")), [])})
    sampled = 0
    for row in rows:
        if sampled >= limit:
            break
        sampled += 1
        for name in relevant:
            idx = indices[name]
            value = "" if idx >= len(row) or row[idx] is None else str(row[idx]).strip()
            if not value or value.lower() in {"nan", "na", "none", "null"}:
                missing[name] += 1
            elif name in invalid and not AA_RE.fullmatch(value):
                invalid[name] += 1
    return sampled, dict(missing), dict(invalid)


def audit_delimited(path: Path, root: Path, sample_limit: int) -> InventoryRow:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    notes: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        prefix = []
        header_idx = None
        for idx, row in enumerate(reader):
            prefix.append(row)
            if looks_like_header(row):
                header_idx = idx
                break
            if idx >= 49:
                break
        if header_idx is None:
            header_idx = 0
            notes.append("header_not_confidently_detected")
        elif header_idx > 0:
            notes.append(f"preamble_rows_skipped:{header_idx}")
        columns = [value.strip() or f"unnamed_{i}" for i, value in enumerate(prefix[header_idx])]

    roles = infer_roles(columns)
    row_count = bad_width = 0
    samples: list[list[str]] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for _ in range(header_idx + 1):
            next(reader, None)
        for row in reader:
            row_count += 1
            bad_width += len(row) != len(columns)
            if len(samples) < sample_limit:
                samples.append(row)
    sampled, missing, invalid = _sample_stats(columns, samples, roles, sample_limit)
    if bad_width:
        notes.append("inconsistent_column_width")
    if not roles["heavy"]:
        notes.append("heavy_or_vhh_column_requires_review")
    if not roles["label"]:
        notes.append("label_column_requires_review")
    return InventoryRow(
        path=path.relative_to(root).as_posix(), format=path.suffix.lower().lstrip("."),
        size_bytes=path.stat().st_size, sha256=sha256_file(path), sheet="",
        header_row=header_idx + 1, row_count=row_count, column_count=len(columns),
        bad_width_rows=bad_width, columns=columns, heavy_candidates=roles["heavy"],
        light_candidates=roles["light"], antigen_candidates=roles["antigen"],
        label_candidates=roles["label"], sampled_rows=sampled, sampled_missing=missing,
        sampled_invalid_sequence=invalid, review_status="needs_human_semantic_review", notes=notes,
    )


def audit_xlsx(path: Path, root: Path, sample_limit: int) -> list[InventoryRow]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    digest = sha256_file(path)
    results = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        columns = [str(value).strip() if value is not None else f"unnamed_{i}" for i, value in enumerate(header)]
        roles = infer_roles(columns)
        samples = []
        for _, row in zip(range(sample_limit), rows):
            samples.append(row)
        sampled, missing, invalid = _sample_stats(columns, samples, roles, sample_limit)
        notes = ["xlsx_row_count_uses_worksheet_dimension"]
        if not roles["heavy"]:
            notes.append("heavy_or_vhh_column_requires_review")
        if not roles["label"]:
            notes.append("label_column_requires_review")
        results.append(InventoryRow(
            path=path.relative_to(root).as_posix(), format="xlsx", size_bytes=path.stat().st_size,
            sha256=digest, sheet=sheet.title, header_row=1,
            row_count=max((sheet.max_row or 1) - 1, 0), column_count=len(columns),
            bad_width_rows=None, columns=columns, heavy_candidates=roles["heavy"],
            light_candidates=roles["light"], antigen_candidates=roles["antigen"],
            label_candidates=roles["label"], sampled_rows=sampled, sampled_missing=missing,
            sampled_invalid_sequence=invalid, review_status="needs_human_semantic_review", notes=notes,
        ))
    workbook.close()
    return results


def write_outputs(rows: list[InventoryRow], output_json: Path, output_csv: Path) -> None:
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema_version": 1,
        "semantics": "Field roles are heuristics and are not approved label registrations.",
        "records": [asdict(row) for row in rows],
    }
    output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fieldnames = list(asdict(rows[0]).keys()) if rows else list(InventoryRow.__annotations__)
    with output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            record = asdict(row)
            for key, value in record.items():
                if isinstance(value, (list, dict)):
                    record[key] = json.dumps(value, ensure_ascii=False, sort_keys=True)
            writer.writerow(record)


def write_report(rows: list[InventoryRow], output: Path) -> None:
    hashes: dict[str, set[str]] = {}
    for row in rows:
        hashes.setdefault(row.sha256, set()).add(row.path)
    duplicate_groups = [sorted(paths) for paths in hashes.values() if len(paths) > 1]
    preambles = [row for row in rows if (row.header_row or 1) > 1]
    bad_width = [row for row in rows if row.bad_width_rows]
    no_heavy = [row for row in rows if not row.heavy_candidates]
    no_label = [row for row in rows if not row.label_candidates]
    total_rows = sum(row.row_count or 0 for row in rows)
    unique_rows = total_rows
    for group in duplicate_groups:
        duplicate_paths = set(group[1:])
        unique_rows -= sum(row.row_count or 0 for row in rows if row.path in duplicate_paths)
    largest = sorted(rows, key=lambda row: row.row_count or 0, reverse=True)[:10]
    lines = [
        "# 阶段1：首轮数据清单与结构审计", "",
        "> 本报告由 `scripts/audit_stage1.py` 生成。字段角色只是启发式候选，尚未完成标签语义人工注册。", "",
        "## 结论", "",
        f"- 审计记录：{len(rows)}（CSV/TSV 每文件一条，XLSX 每工作表一条）。",
        f"- 物理行数合计：{total_rows:,}；扣除已确认的完整重复文件后约 {unique_rows:,} 行。",
        f"- 前导说明行/错位表头：{len(preambles)} 个文件。",
        f"- CSV/TSV 列宽异常：{len(bad_width)} 个文件。",
        f"- 未自动识别候选序列字段：{len(no_heavy)} 条记录；未自动识别标签字段：{len(no_label)} 条记录。",
        "- 所有源文件仅被读取；没有解压 ZIP，也没有向 `data/` 写入。", "",
        "## 已确认异常", "",
    ]
    if preambles:
        lines.extend(f"- `{row.path}`：真实表头在第 {row.header_row} 行。" for row in preambles)
    if duplicate_groups:
        lines.extend(f"- 完整重复文件（SHA-256 相同）：`{'`、`'.join(group)}`。" for group in duplicate_groups)
    if not preambles and not duplicate_groups:
        lines.append("- 无。")
    lines.extend(["", "## 最大数据表", "", "| 行数 | 路径 |", "|---:|---|"])
    lines.extend(f"| {(row.row_count or 0):,} | `{row.path}` |" for row in largest)
    lines.extend(["", "## 当前边界", "",
        "1. `label_candidates` 仅按列名推断，不能替代对论文、单位、方向、截断和 assay 的人工登记。",
        "2. 序列合法性与缺失统计只覆盖每个表前若干样本行，详见 inventory 的 `sampled_rows`。",
        "3. XLSX 行数取工作表 dimension；正式标准化前需逐表复核隐藏行、公式和合并单元格。",
        "4. 结构 ZIP 与 INDI2 ZIP 本轮仍只保留为归档，不进行解压。", "",
        "## 下一验收门", "",
        "逐文件建立标签注册表，至少确认 `study_id / assay_type / unit / direction / label_column / quality_tier / comparable_group / censoring`，未确认文件默认隔离。",
    ])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, default=Path("data"))
    parser.add_argument("--output-json", type=Path, default=Path("data/processed/stage1_inventory.json"))
    parser.add_argument("--output-csv", type=Path, default=Path("data/processed/stage1_inventory.csv"))
    parser.add_argument("--report", type=Path, default=Path("artifacts/stage1_inventory_report.md"))
    parser.add_argument("--sample-rows", type=int, default=10_000)
    args = parser.parse_args()
    root = args.data_root.resolve()
    output_paths = {args.output_json.resolve(), args.output_csv.resolve(), args.report.resolve()}
    if any(root == output or root in output.parents for output in output_paths):
        parser.error("outputs must be outside the read-only data root")
    files = sorted(path for path in root.rglob("*") if path.suffix.lower() in {".csv", ".tsv", ".xlsx"})
    rows: list[InventoryRow] = []
    for path in files:
        if path.suffix.lower() == ".xlsx":
            rows.extend(audit_xlsx(path, root, args.sample_rows))
        else:
            rows.append(audit_delimited(path, root, args.sample_rows))
    write_outputs(rows, args.output_json, args.output_csv)
    write_report(rows, args.report)
    print(json.dumps({"files": len(files), "inventory_records": len(rows), "output": str(args.output_json)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
