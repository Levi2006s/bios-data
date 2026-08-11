#!/usr/bin/env python3
"""Sample numeric label columns and measure their rank relationships."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path

import numpy as np
from scipy.stats import spearmanr


STRICT_NUMBER = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
CENSOR_PREFIX = re.compile(r"^\s*(<=|>=|<|>)")


def parse_number(value: object) -> tuple[float | None, bool]:
    text = "" if value is None else str(value).strip()
    censored = bool(CENSOR_PREFIX.match(text))
    if censored or not STRICT_NUMBER.fullmatch(text):
        return None, censored
    number = float(text)
    return (number if math.isfinite(number) else None), False


def rows_for(record: dict, data_root: Path, limit: int):
    path = data_root / record["path"]
    if record["format"] == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[record["sheet"]]
        iterator = sheet.iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(iterator)]
        yield header
        for index, row in enumerate(iterator):
            if index >= limit:
                break
            yield list(row)
        workbook.close()
        return
    delimiter = "\t" if record["format"] == "tsv" else ","
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for _ in range(record["header_row"] - 1):
            next(reader, None)
        yield next(reader)
        for index, row in enumerate(reader):
            if index >= limit:
                break
            yield row


def profile(record: dict, data_root: Path, limit: int) -> dict:
    iterator = rows_for(record, data_root, limit)
    header = next(iterator)
    labels = record["label_candidates"]
    indices = {name: header.index(name) for name in labels if name in header}
    values = {name: [] for name in indices}
    censored = {name: 0 for name in indices}
    sampled = 0
    for row in iterator:
        sampled += 1
        for name, index in indices.items():
            raw = row[index] if index < len(row) else None
            number, is_censored = parse_number(raw)
            values[name].append(number)
            censored[name] += int(is_censored)
    stats = {}
    for name, column in values.items():
        numeric = np.array([value for value in column if value is not None], dtype=float)
        stats[name] = {
            "numeric_count": int(numeric.size), "missing_or_non_numeric": sampled - int(numeric.size),
            "censored_count": censored[name], "unique_count": int(np.unique(numeric).size),
            "min": float(numeric.min()) if numeric.size else None,
            "max": float(numeric.max()) if numeric.size else None,
        }
    correlations = []
    if "fitness" in values:
        for name in labels:
            if name == "fitness" or name not in values:
                continue
            pairs = [(a, b) for a, b in zip(values["fitness"], values[name]) if a is not None and b is not None]
            rho = float(spearmanr(*zip(*pairs)).statistic) if len(pairs) >= 3 else None
            correlations.append({"other_label": name, "paired_count": len(pairs), "spearman_with_fitness": rho})
    return {"path": record["path"], "sheet": record["sheet"], "sampled_rows": sampled,
            "label_stats": stats, "fitness_correlations": correlations}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inventory", type=Path, default=Path("data/processed/stage1_inventory.json"))
    parser.add_argument("--data-root", type=Path, default=Path("data"))
    parser.add_argument("--output", type=Path, default=Path("data/processed/stage1_label_profiles.json"))
    parser.add_argument("--sample-rows", type=int, default=50_000)
    args = parser.parse_args()
    records = json.loads(args.inventory.read_text(encoding="utf-8"))["records"]
    profiles = [profile(record, args.data_root, args.sample_rows) for record in records]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps({"schema_version": 1, "profiles": profiles}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"profiles": len(profiles), "output": str(args.output)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
